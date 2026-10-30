"""Excel and CSV export manager."""
import csv
import sys
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Optional

sys.path.insert(0, str(Path(__file__).parent.parent))
from config import EXCEL_COLUMNS, EXPORTS_DIR
from utils.logger import get_logger

log = get_logger("export_manager")


class ExportManager:
    """Handles export to XLSX and CSV with rich formatting."""

    def __init__(self, export_dir: Path = EXPORTS_DIR):
        self.export_dir = export_dir
        self.export_dir.mkdir(parents=True, exist_ok=True)

    def _build_filename(self, ext: str) -> Path:
        stamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        return self.export_dir / f"TG_Contacts_{stamp}.{ext}"

    # ─── Excel ──────────────────────────────────────────────────────────────
    def export_excel(self, users: List[Dict], dest_path: Optional[Path] = None) -> Path:
        """Export users to a formatted XLSX file."""
        try:
            import openpyxl
            from openpyxl.styles import (
                Font, PatternFill, Alignment, Border, Side
            )
            from openpyxl.utils import get_column_letter
        except ImportError:
            raise RuntimeError("openpyxl is not installed. Run: pip install openpyxl")

        if dest_path is None:
            dest_path = self._build_filename("xlsx")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "TG Contacts"

        # ── Header style ────────────────────────────────────────────────────
        header_fill = PatternFill("solid", fgColor="1A2035")
        header_font = Font(name="Segoe UI", bold=True, color="F1F5F9", size=11)
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=False)

        thin_side = Side(style="thin", color="1E293B")
        thin_border = Border(left=thin_side, right=thin_side, bottom=thin_side)

        headers = [col[1] for col in EXCEL_COLUMNS]
        keys    = [col[0] for col in EXCEL_COLUMNS]

        ws.append(headers)
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        # Freeze top row
        ws.freeze_panes = "A2"

        # Enable auto-filter
        ws.auto_filter.ref = ws.dimensions

        # ── Data rows ───────────────────────────────────────────────────────
        alt_fill = PatternFill("solid", fgColor="141D2F")
        data_font = Font(name="Segoe UI", size=10, color="F1F5F9")
        alt_font  = Font(name="Segoe UI", size=10, color="94A3B8")

        for row_idx, user in enumerate(users, start=2):
            is_alt = (row_idx % 2 == 0)
            row_fill = alt_fill if is_alt else None
            row_font = alt_font if is_alt else data_font

            for col_idx, key in enumerate(keys, start=1):
                value = user.get(key)

                # Format booleans
                if key == "is_contact":
                    value = "Yes" if value else "No"

                # Null display
                if value is None or value == "":
                    value = ""

                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = row_font
                if row_fill:
                    cell.fill = row_fill

                # Phone as text to preserve leading +
                if key == "phone" and value:
                    cell.number_format = "@"

        # ── Column widths ───────────────────────────────────────────────────
        col_widths = {
            "user_id": 14,
            "first_name": 18,
            "last_name": 18,
            "full_name": 28,
            "username": 22,
            "phone": 20,
            "is_contact": 13,
            "last_message_date": 22,
            "chat_type": 14,
            "extraction_date": 22,
        }
        for col_idx, key in enumerate(keys, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = col_widths.get(key, 18)

        # Row height
        ws.row_dimensions[1].height = 22
        for r in range(2, len(users) + 2):
            ws.row_dimensions[r].height = 18

        wb.save(str(dest_path))
        log.info("Excel exported: %s (%d rows)", dest_path.name, len(users))
        return dest_path

    # ─── CSV ────────────────────────────────────────────────────────────────
    def export_csv(self, users: List[Dict], dest_path: Optional[Path] = None) -> Path:
        """Export users to CSV."""
        if dest_path is None:
            dest_path = self._build_filename("csv")

        headers = [col[1] for col in EXCEL_COLUMNS]
        keys    = [col[0] for col in EXCEL_COLUMNS]

        with open(dest_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=headers)
            writer.writeheader()
            for user in users:
                row = {}
                for key, header in zip(keys, headers):
                    v = user.get(key)
                    if key == "is_contact":
                        v = "Yes" if v else "No"
                    row[header] = v if v is not None else ""
                writer.writerow(row)

        log.info("CSV exported: %s (%d rows)", dest_path.name, len(users))
        return dest_path
